from flask import Flask, render_template, request, send_from_directory
from docx import Document
from fpdf import FPDF

app = Flask(__name__)

# Set the path for uploading files
app.config['UPLOAD_FOLDER'] = 'uploads'

@app.route('/word')
def word():
    return render_template('word.html')

@app.route('/', methods=['POST'])
def contact():
    # Get the uploaded file
    file = request.files['file']

    # Save the file to the uploads folder
    file.save(f"{app.config['UPLOAD_FOLDER']}/{file.filename}")

    # Convert the Word document to PDF
    doc = Document(f"{app.config['UPLOAD_FOLDER']}/{file.filename}")
    pdf = FPDF()
    
    for para in doc.paragraphs:
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(0, 10, txt=para.text, ln=1)

    pdf.output(f"{app.config['UPLOAD_FOLDER']}/{file.filename.split('.')[0]}.pdf")

    return render_template('download_word.html', filename=file.filename)

@app.route('/download/<path:filename>', methods=['GET'])
def download(filename):
    # Serve the converted PDF file for download
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
